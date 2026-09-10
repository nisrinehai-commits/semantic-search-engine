from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "rapport"
SHOTS = OUT / "screenshots"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "Rapport_PFA_INRH_GED_Intelligente_Revu.docx"

BLUE = "1F4E79"
TEAL = "0F766E"
INK = "1A1A18"
MUTED = "66706D"
LINE = "D9E2E0"
LIGHT = "EAF4F2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color=LINE):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)
    tc_pr.append(borders)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(field)
    run._r.append(instr)
    run._r.append(end)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_text(doc, text, style=None, align=None, space_after=8, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.22
    if keep:
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, 10.5, INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, 10.2, INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    set_run_font(run, 17 if level == 1 else 12.5, BLUE if level == 1 else TEAL, bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, 9, MUTED, italic=True)
    return p


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF8E7")
    set_cell_border(cell, "E5C77A")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    first = p.add_run(label + " ")
    set_run_font(first, 9.5, "795A16", bold=True)
    rest = p.add_run(text)
    set_run_font(rest, 9.5, "795A16")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value)
        set_run_font(run, 8.5, "FFFFFF", bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, "F7FAF9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(value))
            set_run_font(run, 8.7, INK)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_figure(doc, filename, caption, width=15.5):
    image = SHOTS / filename
    if image.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together = True
        p.add_run().add_picture(str(image), width=Cm(width))
        add_caption(doc, caption)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.22
    for name, size, color in [("Heading 1", 17, BLUE), ("Heading 2", 12.5, TEAL), ("Heading 3", 11, BLUE)]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    header = section.header.paragraphs[0]
    header.text = "INRH | Rapport de stage PFA | GED intelligente"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, 8, MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INRH")
    set_run_font(run, 26, BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Institut National de Recherche Halieutique")
    set_run_font(run, 12, MUTED)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RAPPORT DE STAGE PFA")
    set_run_font(run, 17, TEAL, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Conception et développement d'une GED intelligente\norientée recherche sémantique")
    set_run_font(run, 22, INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Plateforme documentaire pour l'Institut National de Recherche Halieutique")
    set_run_font(run, 12, MUTED, italic=True)
    for _ in range(4):
        doc.add_paragraph()
    info = [
        ("Présenté par", "Nisrine Haimeur"),
        ("Filière", "Génie Informatique"),
        ("Encadrante professionnelle", "Mme Sara Berguia"),
        ("Service d'accueil", "Service des Systèmes d'Information, de Documentation et Web"),
        ("Période de stage", "du 1er juillet 2026 au 31 août 2026"),
        ("Année universitaire", "2025 - 2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in info:
        cells = table.add_row().cells
        cells[0].width = Cm(5.8)
        cells[1].width = Cm(9.5)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r0 = p0.add_run(label + " :")
        set_run_font(r0, 10.5, MUTED, bold=True)
        p1 = cells[1].paragraphs[0]
        r1 = p1.add_run(value)
        set_run_font(r1, 10.5, INK)
    doc.add_page_break()


def front_matter(doc):
    add_heading(doc, "Remerciements", 1)
    add_text(doc, "Je tiens à remercier Mme Sara Berguia, encadrante professionnelle au sein du Service des Systèmes d'Information, de Documentation et Web de l'INRH, pour son accompagnement, ses retours structurants et la confiance accordée durant ce stage. Ses observations ont permis de faire évoluer le travail d'un moteur de recherche documentaire vers une démonstration de GED intelligente plus cohérente avec les besoins métiers.")
    add_text(doc, "Je remercie également l'ensemble des collaborateurs de l'INRH pour leur accueil et pour les échanges qui ont permis de mieux comprendre la diversité des documents scientifiques, techniques et administratifs produits par l'institut. Enfin, j'adresse mes remerciements à l'équipe pédagogique de la Faculté des Sciences et Techniques de Settat pour les acquis mobilisés dans ce projet.")
    doc.add_page_break()

    add_heading(doc, "Résumé", 1)
    add_text(doc, "Ce rapport présente la conception et la réalisation d'un prototype de Gestion Électronique des Documents (GED) intelligente destiné à l'Institut National de Recherche Halieutique. Le besoin traité consiste à améliorer l'accès à un fonds documentaire hétérogène composé de rapports scientifiques, procédures, études et documents administratifs, pour lesquels une recherche par mots-clés seule reste insuffisante.")
    add_text(doc, "La solution développée met en œuvre une chaîne complète : ingestion de documents PDF, DOCX et TXT, extraction de texte avec prise en charge des PDF scannés par OCR, découpage en passages, génération d'embeddings multilingues SBERT, indexation vectorielle FAISS et recherche hybride combinant proximité sémantique et correspondance lexicale. Cette capacité est intégrée dans une interface GED comprenant une bibliothèque documentaire, des métadonnées, un workflow, une corbeille restaurable, une authentification JWT et une administration par rôles.")
    add_text(doc, "Le prototype propose aussi un assistant documentaire local fondé sur la recherche de passages pertinents et l'affichage de sources citées. L'objectif n'est pas de remplacer une GED d'entreprise complète, mais de démontrer de façon fonctionnelle et explicable la valeur ajoutée de la recherche sémantique dans la valorisation du patrimoine documentaire de l'INRH.")
    add_text(doc, "Mots-clés : GED, recherche sémantique, SBERT, FAISS, FastAPI, OCR, workflow documentaire, JWT, RAG local.")
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_text(doc, "This report presents the design and implementation of an intelligent Electronic Document Management prototype for the National Institute of Fisheries Research. The project addresses access to a heterogeneous collection of scientific, technical and administrative documents, where keyword-only retrieval is often insufficient.")
    add_text(doc, "The solution implements document ingestion for PDF, DOCX and TXT files, text extraction with OCR fallback, chunking, multilingual SBERT embeddings, FAISS vector indexing and hybrid retrieval. The retrieval engine is integrated into a GED interface with document metadata, workflow, recoverable deletion, JWT authentication, role-based administration and a local document assistant with cited sources.")
    add_text(doc, "Keywords: electronic document management, semantic search, SBERT, FAISS, FastAPI, OCR, document workflow, JWT, local RAG.")
    doc.add_page_break()

    add_heading(doc, "Liste des abréviations", 1)
    add_table(doc, ["Abréviation", "Signification"], [
        ("INRH", "Institut National de Recherche Halieutique"),
        ("GED", "Gestion Électronique des Documents"),
        ("SSID", "Service des Systèmes d'Information, de Documentation et Web"),
        ("NLP", "Natural Language Processing"),
        ("OCR", "Reconnaissance optique de caractères"),
        ("SBERT", "Sentence-BERT"),
        ("FAISS", "Facebook AI Similarity Search"),
        ("JWT", "JSON Web Token"),
        ("API", "Application Programming Interface"),
        ("RAG", "Retrieval-Augmented Generation"),
    ], [3.0, 13.0])
    doc.add_page_break()

    add_heading(doc, "Sommaire", 1)
    for line in [
        "Introduction générale",
        "Chapitre 1 - Contexte général et problématique",
        "Chapitre 2 - Étude de l'existant et fondements théoriques",
        "Chapitre 3 - Analyse des besoins et conception de la solution",
        "Chapitre 4 - Réalisation de la GED intelligente",
        "Chapitre 5 - Validation, résultats et limites",
        "Conclusion générale et perspectives",
        "Bibliographie",
        "Annexe - Principales routes de l'API",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(line)
        set_run_font(run, 11, INK)
    doc.add_page_break()


def chapter_one(doc):
    add_heading(doc, "Introduction générale", 1)
    add_text(doc, "Les organisations produisent continuellement des documents dont la valeur ne dépend pas seulement de leur conservation, mais aussi de leur capacité à être retrouvés, compris et réutilisés. À l'INRH, les rapports de campagnes, études de ressources halieutiques, procédures de laboratoire, conventions et documents de qualité constituent une mémoire scientifique et opérationnelle importante. Leur hétérogénéité rend toutefois la recherche difficile lorsque l'utilisateur ne connaît pas les termes exacts employés dans un document.")
    add_text(doc, "Le présent travail répond à cette difficulté par la réalisation d'un prototype de GED intelligente. La contribution principale est l'intégration d'une recherche sémantique capable de rapprocher une requête du sens des passages documentaires, complétée par une recherche lexicale afin de conserver la précision des termes métier. Le rapport expose la démarche de conception, l'architecture, les écrans réalisés, les tests et les limites à prendre en compte avant une mise en production.")
    doc.add_page_break()

    add_heading(doc, "Chapitre 1  Contexte général et problématique", 1)
    add_heading(doc, "1.1 Organisme d'accueil et service concerné", 2)
    add_text(doc, "L'Institut National de Recherche Halieutique est un établissement public marocain à caractère scientifique et technique. Ses missions couvrent notamment l'évaluation des ressources halieutiques, l'appui au développement de l'aquaculture, la conduite de campagnes océanographiques et la production d'avis scientifiques. Ces activités génèrent un volume important de documents scientifiques, techniques et administratifs.")
    add_text(doc, "Le stage a été réalisé au sein du Service des Systèmes d'Information, de Documentation et Web. Selon les éléments communiqués dans le rapport initial, ce service recueille les besoins des directions métier, contribue à l'intégration des solutions applicatives, assure la gestion des systèmes d'information, développe des outils internes, centralise le fonds documentaire et veille à la sécurité et à l'intégrité des données. Cette mission de gouvernance documentaire justifie directement le sujet étudié.")
    add_heading(doc, "1.2 Contexte et problématique", 2)
    add_text(doc, "Une recherche fondée exclusivement sur les mots-clés suppose que la formulation de l'utilisateur recouvre celle du document. Or, dans un contexte scientifique, une même notion peut être exprimée par des formulations différentes : ressources halieutiques, stocks de pêche, biomasse exploitable ou suivi des espèces. Le document pertinent peut donc rester invisible malgré l'existence d'un contenu sémantiquement proche.")
    add_note(doc, "Problématique", "Comment concevoir une GED qui permet de retrouver des passages pertinents dans des documents hétérogènes, tout en assurant une gestion documentaire, une traçabilité et des contrôles d'accès adaptés à une organisation ?")
    add_heading(doc, "1.3 Objectifs", 2)
    for text in [
        "Construire une chaîne d'ingestion compatible avec les formats PDF, DOCX et TXT, avec bascule OCR pour les PDF dont le texte natif est insuffisant.",
        "Mettre en œuvre une recherche hybride associant recherche sémantique par embeddings et recherche lexicale.",
        "Intégrer le moteur dans une interface GED : bibliothèque, métadonnées, prévisualisation, téléchargement et cycle de vie documentaire.",
        "Proposer des mécanismes de sécurité et de gouvernance : authentification JWT, rôles, permissions, journal d'activité, archivage et corbeille restaurable.",
        "Démontrer les usages de l'IA documentaire au travers d'un assistant fondé sur les passages retrouvés et leurs sources.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "1.4 Démarche adoptée", 2)
    add_text(doc, "La démarche a été incrémentale : étude des approches de recherche, conception, développement du pipeline documentaire, intégration des fonctions de gestion et validation des routes principales. Cette progression a permis de faire évoluer l'interface depuis un prototype technique vers une démonstration métier plus lisible.")


def chapter_two(doc):
    doc.add_page_break()
    add_heading(doc, "Chapitre 2  Étude de l'existant et fondements théoriques", 1)
    add_heading(doc, "2.1 Exigences documentaires et solutions GED", 2)
    add_text(doc, "Une GED professionnelle doit assurer l'identification du document, son accès contrôlé, sa conservation, sa traçabilité et la maîtrise de ses modifications. Dans une logique inspirée de l'ISO 9001, l'information documentée doit être disponible au bon moment, protégée contre les accès ou modifications non autorisés et maintenue lisible durant sa période de conservation.")
    add_table(doc, ["Fonction attendue", "Réponse apportée par le prototype", "Limite ou évolution"], [
        ("Centralisation", "Bibliothèque et import de PDF, DOCX, TXT", "Stockage objet ou réseau à prévoir en production"),
        ("Accès", "JWT, rôles et permissions par document", "Connexion LDAP ou Active Directory à intégrer"),
        ("Traçabilité", "Workflow, historique et journal d'activité", "Journal d'audit persistant à renforcer"),
        ("Conservation", "Archivage et suppression logique restaurable", "Politique de rétention à formaliser"),
        ("Recherche", "Recherche hybride et passages pertinents", "Évaluation sur corpus INRH réel à poursuivre"),
    ], [3.3, 7.0, 6.0])
    add_heading(doc, "2.2 Recherche lexicale et ses limites", 2)
    add_text(doc, "La recherche lexicale compare principalement les termes de la requête à ceux des documents. Des méthodes comme TF-IDF pondèrent les mots rares et offrent une base rapide et interprétable. Elles restent toutefois sensibles aux variantes, aux synonymes et aux différences de formulation. Elles sont donc conservées dans le prototype, mais ne constituent pas l'unique mécanisme de classement.")
    add_heading(doc, "2.3 Recherche sémantique avec SBERT", 2)
    add_text(doc, "SBERT représente une phrase ou un passage sous la forme d'un vecteur dense, appelé embedding. Des textes proches par le sens ont des vecteurs proches dans l'espace vectoriel, même lorsqu'ils ne partagent pas les mêmes mots. Le projet s'appuie sur un modèle multilingue compatible avec le français : paraphrase-multilingual-MiniLM-L12-v2. Cette option est adaptée à un prototype local, car elle réduit la dépendance à un service externe et permet de traiter des documents potentiellement sensibles.")
    add_heading(doc, "2.4 FAISS et recherche hybride", 2)
    add_text(doc, "FAISS permet d'indexer les vecteurs et de retrouver rapidement les passages les plus proches d'une requête. Les embeddings sont normalisés et comparés par produit scalaire, ce qui correspond à une similarité cosinus après normalisation. Le moteur fusionne ensuite le score sémantique et le score lexical. Dans la configuration par défaut, le score final est calculé selon : score final = 0,7 x score sémantique + 0,3 x score lexical.")
    add_text(doc, "Le choix d'une recherche hybride répond à un compromis : la composante sémantique améliore le rappel lorsque le vocabulaire varie, tandis que la composante lexicale conserve l'importance des termes techniques, sigles et noms propres. Le poids de chaque composante reste configurable au niveau de la requête API.")


def chapter_three(doc):
    doc.add_page_break()
    add_heading(doc, "Chapitre 3  Analyse des besoins et conception de la solution", 1)
    add_heading(doc, "3.1 Acteurs et besoins", 2)
    add_table(doc, ["Acteur", "Besoins principaux", "Droits dans le prototype"], [
        ("Employé", "Importer, consulter, rechercher, questionner l'assistant", "Lecture, dépôt et soumission"),
        ("Validateur", "Contrôler un document soumis et commenter son état", "Lecture et validation"),
        ("Responsable qualité", "Gérer les métadonnées et le cycle documentaire", "Écriture, validation, workflow"),
        ("Administrateur", "Piloter les utilisateurs, les archives et la corbeille", "Accès complet et gouvernance"),
    ], [3.0, 8.1, 5.2])
    add_heading(doc, "3.2 Cas d'utilisation prioritaires", 2)
    for text in [
        "Importer un document, vérifier son format, extraire son texte et reconstruire l'index de recherche.",
        "Rechercher un document par titre, catégorie, mots-clés ou intention exprimée en langage naturel.",
        "Prévisualiser un résultat directement depuis la page de recherche avant d'ouvrir sa fiche complète.",
        "Interroger l'assistant documentaire et consulter les sources mobilisées dans la réponse.",
        "Modifier les métadonnées, faire progresser un document dans le workflow, l'archiver ou le restaurer depuis la corbeille selon le rôle.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "3.3 Architecture fonctionnelle", 2)
    add_table(doc, ["Couche", "Responsabilité", "Technologies"], [
        ("Interface", "Bibliothèque, recherche, assistant, administration", "HTML, CSS, JavaScript"),
        ("API REST", "Authentification, documents, workflow, recherche", "FastAPI, Pydantic"),
        ("Ingestion", "Extraction PDF, DOCX, TXT et OCR", "pypdf, python-docx, pytesseract, pdf2image"),
        ("Recherche", "Embeddings, index FAISS, fusion hybride", "sentence-transformers, FAISS, scikit-learn"),
        ("Persistance", "Métadonnées, actions, versions, utilisateurs", "SQLite par défaut, PostgreSQL configurable"),
    ], [3.0, 7.3, 6.0])
    add_heading(doc, "3.4 Pipeline documentaire", 2)
    add_text(doc, "Le pipeline suit la séquence suivante : dépôt du fichier, contrôle de l'extension, stockage de l'original, extraction du texte, bascule OCR si nécessaire, nettoyage minimal, découpage en passages avec chevauchement, génération des embeddings, construction ou mise à jour de l'index FAISS, puis exposition des résultats via l'API. La recherche est donc effectuée au niveau du passage, ce qui permet d'afficher l'extrait réellement utile plutôt que de renvoyer uniquement le titre d'un document entier.")
    add_heading(doc, "3.5 Modèle de données et cycle de vie", 2)
    add_text(doc, "Les métadonnées couvrent notamment le titre, l'auteur, la catégorie, la description, les tags, le statut, le dossier et les rôles autorisés. Le cycle documentaire simplifié comprend les états Brouillon, Soumis, En validation, Approuvé, Archivé et Rejeté. La suppression est logique : le fichier n'est pas détruit immédiatement, un indicateur deleted et des informations de date et d'auteur sont conservés, ce qui rend la restauration possible.")
    add_heading(doc, "3.6 Sécurité", 2)
    add_text(doc, "L'API produit un jeton JWT lors de la connexion. Les routes de modification, d'import, de suppression, de restauration et de gestion des utilisateurs vérifient les permissions nécessaires. L'interface adapte également les actions visibles au rôle : l'employé consulte et dépose, alors que l'administration de la gouvernance documentaire est réservée à l'administrateur. Cette séparation améliore l'expérience utilisateur, mais la règle de sécurité reste portée par le backend.")


def chapter_four(doc):
    doc.add_page_break()
    add_heading(doc, "Chapitre 4  Réalisation de la GED intelligente", 1)
    add_heading(doc, "4.1 Organisation du projet", 2)
    add_text(doc, "Le projet est organisé par responsabilités : ingestion, embeddings, recherche, API, configuration et interface statique. Cette structure permet de faire évoluer chaque composant sans mélanger l'extraction documentaire, la logique de recherche et la présentation utilisateur.")
    add_table(doc, ["Répertoire", "Contenu"], [
        ("src/ingestion", "Extraction PDF, DOCX, TXT et OCR"),
        ("src/embeddings", "Découpage, génération des embeddings et indexation"),
        ("src/search", "Recherche sémantique, lexicale et hybride"),
        ("src/api", "Routes REST, schémas, sécurité JWT et persistance"),
        ("static", "Interface web de démonstration"),
        ("tests", "Tests de l'API, de l'ingestion et du moteur"),
    ], [4.0, 12.3])
    add_heading(doc, "4.2 Tableau de bord et bibliothèque documentaire", 2)
    add_text(doc, "Le tableau de bord met en avant les trois actions principales : ajouter un document, retrouver un document et poser une question à l'assistant. Il synthétise les documents disponibles, en attente et archivés, les derniers ajouts, l'activité récente et la répartition par catégories. La bibliothèque affiche les informations utiles à la décision : type de fichier, auteur, catégorie, statut et date. Les filtres par type, catégorie et texte sont appliqués à la liste.")
    add_figure(doc, "01-dashboard.png", "Figure 1 - Tableau de bord de la GED intelligente INRH", 15.8)
    add_figure(doc, "02-bibliotheque.png", "Figure 2 - Bibliothèque documentaire avec filtres et actions rapides", 15.8)
    add_heading(doc, "4.3 Recherche et prévisualisation", 2)
    add_text(doc, "La page Recherche adopte une interaction proche d'un moteur de recherche : saisie d'une requête, résultats classés, extrait pertinent, catégorie et score. Un aperçu est accessible directement dans une fenêtre modale, sans obliger l'utilisateur à quitter sa recherche. Cette décision améliore la fluidité : l'utilisateur peut confirmer la pertinence d'un document avant d'ouvrir sa fiche complète.")
    add_note(doc, "Transparence de démonstration", "Les données fictives de l'interface sont identifiées comme telles. Elles ne simulent pas un vrai fichier PDF : leur export produit une fiche texte. Un fichier réellement importé conserve son format original et un PDF importé peut être affiché dans un lecteur intégré.")
    add_figure(doc, "03-recherche-apercu.png", "Figure 3 - Résultats de recherche avec aperçu direct d'un document", 15.8)
    add_heading(doc, "4.4 Assistant documentaire", 2)
    add_text(doc, "L'assistant documentaire repose sur la recherche hybride. Une question est transformée en requête, les passages les plus pertinents sont retrouvés, puis une réponse extractive est formée à partir de ces passages. Les sources affichent le document et l'extrait associés. Cette approche est utile pour une démonstration RAG locale, car elle rend la réponse contrôlable et limite les affirmations non sourcées.")
    add_figure(doc, "04-assistant-ia.png", "Figure 4 - Assistant documentaire avec réponse et sources citées", 15.8)
    add_heading(doc, "4.5 Administration et gouvernance", 2)
    add_text(doc, "L'espace Administration est séparé de l'espace employé. Il permet de visualiser les rôles, les utilisateurs, le journal d'activité et surtout la gouvernance documentaire. L'administrateur dispose d'une recherche sur les documents actifs, les documents archivés, la corbeille restaurable et l'ensemble du corpus. Il peut archiver, réactiver ou restaurer un document supprimé logiquement. La conservation du document dans la corbeille répond à un besoin de sécurité opérationnelle et de traçabilité.")
    add_figure(doc, "05-administration.png", "Figure 5 - Espace d'administration et gouvernance documentaire", 15.8)


def chapter_five(doc):
    doc.add_page_break()
    add_heading(doc, "Chapitre 5  Validation, résultats et limites", 1)
    add_heading(doc, "5.1 Stratégie de validation", 2)
    add_text(doc, "La validation combine des tests automatisés et des vérifications fonctionnelles. Les tests couvrent notamment la recherche hybride, les paramètres de recherche, l'extraction de documents, les routes de prévisualisation et téléchargement, l'authentification, les permissions par rôle, les métadonnées, le workflow et la restauration. L'objectif est de vérifier que les fonctionnalités exposées par l'interface correspondent aux règles appliquées côté API.")
    add_table(doc, ["Fonction", "Vérification réalisée", "Résultat attendu"], [
        ("Ingestion", "Contrôle des extensions PDF, DOCX, TXT", "Formats acceptés et stockage contrôlé"),
        ("Extraction", "PDF texte, DOCX, TXT et bascule OCR", "Texte exploitable pour indexation"),
        ("Recherche", "Modes semantic, keyword et hybrid", "Résultats classés avec score final"),
        ("Sécurité", "Connexion et permissions par rôle", "Refus des actions non autorisées"),
        ("Cycle de vie", "Mise à jour, suppression logique, restauration", "Traçabilité et récupération possible"),
        ("Interface", "Dashboard, bibliothèque, recherche, assistant, administration", "Parcours cohérent et états visibles"),
    ], [3.1, 8.1, 5.1])
    add_heading(doc, "5.2 Résultats observés", 2)
    add_text(doc, "Le prototype permet de réaliser un parcours complet, depuis l'import jusqu'à la recherche et à la gouvernance. Les résultats de recherche affichent des extraits au niveau du passage, ce qui aide l'utilisateur à comprendre pourquoi un document est proposé. L'assistant ajoute une couche d'interrogation en langage naturel tout en gardant les sources visibles. La séparation entre espace employé et administration rend également les responsabilités plus lisibles.")
    add_text(doc, "Les mesures quantitatives détaillées doivent être interprétées avec prudence tant qu'un corpus représentatif de documents INRH et une vérité terrain validée ne sont pas disponibles. Le rapport ne présente donc pas de chiffres de précision ou de rappel non reproductibles. Une évaluation formelle devra construire un jeu de requêtes, faire annoter les résultats pertinents par les utilisateurs métier, puis comparer les modes lexical, sémantique et hybride sur les mêmes cas.")
    add_heading(doc, "5.3 Difficultés rencontrées et réponses apportées", 2)
    add_table(doc, ["Difficulté", "Réponse apportée"], [
        ("PDF scannés sans texte exploitable", "Détection d'un texte natif insuffisant puis bascule vers l'OCR."),
        ("Variabilité du vocabulaire métier", "Recherche hybride et modèle d'embeddings multilingue."),
        ("Risque de confusion entre données démo et fichiers réels", "Avertissement explicite et export de fiche texte pour les données fictives."),
        ("Besoin de gouvernance après suppression", "Suppression logique, corbeille restaurable, date et acteur de l'action."),
        ("Différence entre les usages métier", "Navigation et actions adaptées au rôle connecté."),
    ], [5.4, 10.9])
    add_heading(doc, "5.4 Limites actuelles", 2)
    for text in [
        "L'interface de démonstration contient un jeu de documents fictifs pour rendre les écrans immédiatement présentables. Le branchement complet de tous les écrans sur un corpus réel doit être finalisé avant exploitation opérationnelle.",
        "SQLite est adapté au développement et à la démonstration. PostgreSQL est prévu comme cible de déploiement par la couche de persistance, mais les migrations et l'administration de production restent à industrialiser.",
        "L'assistant est extractif et local : il cite les passages retrouvés, mais ne remplace pas un modèle de génération métier validé ni un processus de vérification humaine.",
        "Les règles de rétention, la classification de confidentialité et la gestion fine des dossiers doivent être définies avec les responsables métier avant mise en production.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "5.5 Perspectives d'évolution", 2)
    add_table(doc, ["Priorité", "Évolution", "Apport"], [
        ("1", "Connexion LDAP ou Active Directory", "Gestion centralisée des identités INRH"),
        ("1", "Dossiers et permissions héritées", "Gouvernance par service ou domaine métier"),
        ("2", "Migrations PostgreSQL et sauvegardes", "Robustesse et exploitation multi-utilisateur"),
        ("2", "Workflow multi-niveaux paramétrable", "Validation adaptée aux types documentaires"),
        ("3", "Surlignage du passage retrouvé dans les PDF", "Meilleure explicabilité de la recherche"),
        ("3", "Knowledge graph et recommandations", "Navigation transversale dans la connaissance"),
    ], [1.8, 7.3, 7.2])


def ending(doc):
    doc.add_page_break()
    add_heading(doc, "Conclusion générale", 1)
    add_text(doc, "Le travail réalisé a permis de dépasser le cadre d'un moteur de recherche isolé pour produire une démonstration cohérente de GED intelligente. La chaîne technique couvre l'ingestion de formats hétérogènes, l'extraction de texte, l'OCR, le découpage en passages, les embeddings SBERT, l'indexation FAISS et la recherche hybride. Ces composants sont exposés par une API FastAPI et intégrés à une interface conçue autour des usages documentaires prioritaires.")
    add_text(doc, "La valeur du prototype réside dans l'articulation entre recherche sémantique et gouvernance documentaire. L'utilisateur peut retrouver un document malgré les variations de vocabulaire, vérifier un extrait, interroger l'assistant et consulter les sources. L'administrateur peut quant à lui suivre le cycle de vie documentaire, archiver, restaurer depuis la corbeille et visualiser les actions importantes. Cette base constitue un socle crédible pour poursuivre l'industrialisation avec un corpus réel, des règles de gestion validées et une intégration au système d'information de l'INRH.")
    doc.add_page_break()
    add_heading(doc, "Bibliographie", 1)
    refs = [
        "ISO. ISO 9001:2015 - Systèmes de management de la qualité - Exigences.",
        "Reimers, N. et Gurevych, I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP, 2019.",
        "Johnson, J., Douze, M. et Jégou, H. Billion-scale similarity search with GPUs. IEEE Transactions on Big Data, 2019.",
        "FastAPI. Documentation officielle. https://fastapi.tiangolo.com/.",
        "Facebook Research. FAISS documentation. https://faiss.ai/.",
        "Sentence Transformers. Documentation officielle. https://www.sbert.net/.",
        "PostgreSQL Global Development Group. PostgreSQL Documentation. https://www.postgresql.org/docs/.",
    ]
    for ref in refs:
        add_text(doc, ref, space_after=5)
    doc.add_page_break()
    add_heading(doc, "Annexe  Principales routes de l'API", 1)
    add_table(doc, ["Méthode", "Route", "Rôle"], [
        ("POST", "/auth/login", "Connexion et émission du jeton JWT"),
        ("POST", "/upload", "Import du fichier et reconstruction de l'index"),
        ("GET", "/documents", "Liste filtrée et paginée des documents"),
        ("GET", "/documents/{filename}/preview", "Prévisualisation adaptée au type de fichier"),
        ("PUT", "/documents/{filename}/metadata", "Mise à jour des métadonnées"),
        ("POST", "/search", "Recherche lexical, sémantique ou hybride"),
        ("POST", "/assistant/ask", "Assistant documentaire avec sources"),
        ("POST", "/documents/{filename}/workflow", "Transition de workflow et historisation"),
        ("DELETE", "/documents/{filename}", "Suppression logique"),
        ("POST", "/documents/{filename}/restore", "Restauration depuis la corbeille"),
        ("GET", "/users", "Liste des utilisateurs autorisés"),
    ], [2.0, 6.6, 7.7])
    add_text(doc, "Les routes sensibles exigent une authentification JWT et vérifient les permissions associées au rôle de l'utilisateur. La documentation interactive de l'API est également disponible via l'interface OpenAPI de FastAPI lors de l'exécution locale.")


def build():
    doc = Document()
    configure(doc)
    cover(doc)
    front_matter(doc)
    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    ending(doc)
    doc.core_properties.title = "Rapport de stage PFA GED intelligente INRH"
    doc.core_properties.author = "Nisrine Haimeur"
    doc.core_properties.subject = "Conception et développement d'une GED intelligente orientée recherche sémantique"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
