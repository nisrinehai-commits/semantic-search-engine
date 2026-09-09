"""
API REST du moteur de recherche semantique.
"""

import json
import base64
import hmac
import os
import re
from hashlib import sha256
from datetime import datetime
from math import ceil
from pathlib import Path

from fastapi import Depends, FastAPI, Header, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from src.api.database import (
    connect,
    decode_json,
    encode_json,
    execute,
    init_db,
    placeholder,
    rows_to_dicts,
)
from src.api.schemas import (
    DashboardBucket,
    DashboardRecentDocument,
    DashboardResponse,
    AssistantAnswerResponse,
    AssistantQuestionRequest,
    AssistantSourceItem,
    DocumentItem,
    DocumentAIAnalysisResponse,
    DocumentListResponse,
    DocumentMetadata,
    DocumentMetadataResponse,
    DocumentPreviewResponse,
    DocumentVersionItem,
    DocumentVersionsResponse,
    LoginRequest,
    LoginResponse,
    SearchRequest,
    SearchResponse,
    SearchResult,
    SimilarDocumentItem,
    SimilarDocumentsResponse,
    UserItem,
    WorkflowActionItem,
    WorkflowHistoryResponse,
    WorkflowTransitionRequest,
)

EMBEDDINGS_DIR = "data/embeddings"
RAW_DIR = "data/raw"
PROCESSED_DIR = "data/processed"
METADATA_PATH = "data/metadata/document_metadata.json"
WORKFLOW_LOG_PATH = "data/metadata/workflow_history.json"
VERSIONS_PATH = "data/metadata/document_versions.json"
USERS_PATH = "data/metadata/users.json"
SUPPORTED_UPLOAD_EXTENSIONS = {".pdf", ".docx", ".txt"}
WORKFLOW_STEPS = ["Brouillon", "Soumis", "En validation", "Approuve", "Rejete", "Archive"]
JWT_SECRET = os.getenv("GED_JWT_SECRET", "change-this-secret-for-production")
JWT_EXPIRES_SECONDS = int(os.getenv("GED_JWT_EXPIRES_SECONDS", "86400"))
PUBLIC_READ_USER = {
    "id": "anonymous",
    "name": "Visiteur local",
    "email": "anonymous@local",
    "role": "Employe",
    "permissions": ["documents:read", "search:read"],
}
FRENCH_STOPWORDS = {
    "avec", "dans", "des", "du", "elle", "est", "les", "leur", "leurs", "nous",
    "par", "pas", "pour", "que", "qui", "sur", "une", "vous", "aux", "ces",
    "cette", "comme", "document", "documents", "rapport", "entre", "plus",
    "afin", "ainsi", "sont", "etre", "son", "ses", "the", "and", "for",
}
CATEGORY_RULES = {
    "Peche": {"peche", "halieutique", "ressources", "maritime", "poisson"},
    "Aquaculture": {"aquaculture", "production", "elevage", "marine", "larves"},
    "Qualite": {"qualite", "norme", "validation", "procedure", "audit"},
    "RH": {"rh", "ressources", "humaines", "salaire", "conge", "employe"},
    "Finance": {"finance", "budget", "facture", "paiement", "cout", "achat"},
    "Informatique": {"informatique", "systeme", "donnees", "securite", "serveur"},
}

app = FastAPI(
    title="Prototype de recherche semantique documentaire - INRH",
    description="API de recherche semantique par passages, basee sur SBERT et FAISS.",
    version="1.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:8000",
        "http://127.0.0.1:8001",
        "http://127.0.0.1:8010",
        "http://127.0.0.1:8011",
        "http://127.0.0.1:8012",
        "http://127.0.0.1:8013",
        "http://127.0.0.1:8014",
        "http://127.0.0.1:8015",
        "http://127.0.0.1:8016",
        "http://127.0.0.1:8017",
        "http://127.0.0.1:8018",
        "http://127.0.0.1:8019",
        "http://localhost:8000",
        "http://localhost:8015",
        "http://localhost:8016",
        "http://localhost:8017",
        "http://localhost:8018",
        "http://localhost:8019",
        "null",
    ],
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/app", StaticFiles(directory="static", html=True), name="static")


@app.on_event("startup")
def startup():
    init_db(METADATA_PATH)
    migrate_legacy_json_files()
    seed_default_users()


def require_permission(permission: str):
    def dependency(authorization: str | None = Header(default=None)) -> dict:
        user = get_current_user(authorization)
        if not has_permission(user, permission):
            raise HTTPException(status_code=403, detail="Permission insuffisante.")
        return user

    return dependency


def optional_current_user(authorization: str | None = Header(default=None)) -> dict:
    if authorization and authorization.lower().startswith("bearer "):
        return get_current_user(authorization)
    return PUBLIC_READ_USER


def get_current_user(authorization: str | None) -> dict:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Token JWT manquant.")
    payload = decode_jwt(authorization.split(" ", 1)[1])
    users = [user for user in load_users_store() if user["email"].lower() == payload.get("sub", "").lower()]
    if not users:
        raise HTTPException(status_code=401, detail="Utilisateur invalide.")
    return users[0]


def has_permission(user: dict, permission: str) -> bool:
    permissions = set(user.get("permissions", []))
    domain = permission.split(":", 1)[0]
    return permission in permissions or f"{domain}:*" in permissions or "*:*" in permissions


def ensure_document_access(filename: str, user: dict) -> None:
    metadata = get_metadata_for(filename)
    allowed_roles = metadata.get("allowed_roles", [])
    if not allowed_roles:
        return
    if user.get("role") in allowed_roles or has_permission(user, "documents:*"):
        return
    raise HTTPException(status_code=403, detail="Acces interdit pour ce document.")


def encode_segment(data: dict) -> str:
    raw = json.dumps(data, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    return base64.urlsafe_b64encode(raw).rstrip(b"=").decode("ascii")


def decode_segment(segment: str) -> dict:
    padding = "=" * (-len(segment) % 4)
    return json.loads(base64.urlsafe_b64decode(f"{segment}{padding}").decode("utf-8"))


def create_jwt(user: dict) -> str:
    import time

    header = {"alg": "HS256", "typ": "JWT"}
    payload = {
        "sub": user["email"],
        "role": user["role"],
        "iat": int(time.time()),
        "exp": int(time.time()) + JWT_EXPIRES_SECONDS,
    }
    signing_input = f"{encode_segment(header)}.{encode_segment(payload)}"
    signature = hmac.new(JWT_SECRET.encode("utf-8"), signing_input.encode("ascii"), "sha256").digest()
    encoded_signature = base64.urlsafe_b64encode(signature).rstrip(b"=").decode("ascii")
    return f"{signing_input}.{encoded_signature}"


def decode_jwt(token: str) -> dict:
    import time

    try:
        header_segment, payload_segment, signature_segment = token.split(".")
        signing_input = f"{header_segment}.{payload_segment}"
        expected = hmac.new(JWT_SECRET.encode("utf-8"), signing_input.encode("ascii"), "sha256").digest()
        expected_segment = base64.urlsafe_b64encode(expected).rstrip(b"=").decode("ascii")
        if not hmac.compare_digest(expected_segment, signature_segment):
            raise ValueError("signature")
        payload = decode_segment(payload_segment)
        if int(payload.get("exp", 0)) < int(time.time()):
            raise ValueError("expired")
        return payload
    except Exception as exc:
        raise HTTPException(status_code=401, detail="Token JWT invalide.") from exc


@app.get("/")
def root():
    """Message de bienvenue / verification que l'API tourne."""
    return {"message": "Prototype de recherche semantique INRH - API operationnelle."}


@app.get("/health")
def health():
    """Endpoint de sante, utile pour du monitoring futur."""
    return {"status": "ok"}


@app.get("/stats")
def stats():
    """Retourne les statistiques de l'index actuellement disponible."""
    stats_path = Path(EMBEDDINGS_DIR) / "stats.json"
    metadata_path = Path(EMBEDDINGS_DIR) / "metadata.json"

    if not metadata_path.exists():
        raise HTTPException(
            status_code=503,
            detail="Index introuvable. Lance python -m src.embeddings.build_index",
        )

    if stats_path.exists():
        return json.loads(stats_path.read_text(encoding="utf-8"))

    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    return {
        "documents": len({item["filename"] for item in metadata}),
        "chunks": len(metadata),
        "dimension": None,
        "index_type": "FAISS",
    }


@app.get("/dashboard", response_model=DashboardResponse)
def dashboard(current_user: dict = Depends(optional_current_user)):
    """Retourne les indicateurs principaux de la GED."""
    documents = discover_documents(current_user=current_user)
    total_size = sum(document.size_bytes for document in documents)
    workflow_actions = collect_recent_workflow_actions(limit=8)

    return DashboardResponse(
        total_documents=len(documents),
        total_size_bytes=total_size,
        total_size_label=format_size(total_size),
        indexed_documents=sum(1 for document in documents if document.processed),
        pending_documents=sum(
            1 for document in documents
            if document.status in {"Brouillon", "Soumis", "En validation", "Rejete"}
        ),
        by_type=count_buckets(document.type for document in documents),
        by_category=count_buckets(document.category for document in documents),
        by_status=count_buckets(document.status for document in documents),
        recent_documents=[
            DashboardRecentDocument(
                name=document.name,
                title=document.title,
                type=document.type,
                category=document.category,
                status=document.status,
                modified_at=document.modified_at,
            )
            for document in sort_documents(documents, "date_desc")[:6]
        ],
        recent_activity=workflow_actions,
    )


@app.post("/search", response_model=SearchResponse)
def search_documents(request: SearchRequest):
    """Recherche les passages les plus pertinents pour une requete donnee."""
    from src.search.search_engine import search

    try:
        raw_results = search(
            request.query,
            top_k=request.top_k,
            mode=request.mode,
            semantic_weight=request.semantic_weight,
            keyword_weight=request.keyword_weight,
        )
    except FileNotFoundError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    results = [SearchResult(**result) for result in raw_results]
    return SearchResponse(query=request.query, results=results)


@app.get("/documents", response_model=DocumentListResponse)
def list_documents(
    q: str = "",
    file_type: str = "all",
    category: str = "all",
    sort_by: str = "date_desc",
    include_deleted: bool = False,
    page: int = Query(1, ge=1),
    per_page: int = Query(8, ge=1, le=50),
    current_user: dict = Depends(optional_current_user),
):
    """Liste les documents importes avec recherche, filtre, tri et pagination."""
    documents = discover_documents(include_deleted=include_deleted, current_user=current_user)

    if q.strip():
        needle = q.strip().lower()
        documents = [
            doc for doc in documents
            if (
                needle in doc.title.lower()
                or needle in doc.name.lower()
                or needle in doc.description.lower()
                or needle in doc.author.lower()
                or any(needle in tag.lower() for tag in doc.tags)
                or needle in doc.category.lower()
            )
        ]

    if file_type != "all":
        documents = [doc for doc in documents if doc.type.lower() == file_type.lower()]

    if category != "all":
        documents = [doc for doc in documents if doc.category.lower() == category.lower()]

    documents = sort_documents(documents, sort_by)
    total = len(documents)
    pages = max(1, ceil(total / per_page))
    start = (page - 1) * per_page
    end = start + per_page

    return DocumentListResponse(
        items=documents[start:end],
        total=total,
        page=page,
        per_page=per_page,
        pages=pages,
    )


@app.get("/documents/{filename}/download")
def download_document(filename: str, current_user: dict = Depends(optional_current_user)):
    """Telecharge le document original de maniere controlee."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    return FileResponse(path, filename=path.name)


@app.get("/documents/{filename}/preview", response_model=DocumentPreviewResponse)
def preview_document(filename: str, current_user: dict = Depends(optional_current_user)):
    """Retourne une previsualisation textuelle ou une URL integree pour PDF."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return DocumentPreviewResponse(
            name=path.name,
            type="PDF",
            mode="embed",
            content=f"/documents/{path.name}/download",
        )

    if extension == ".txt":
        return DocumentPreviewResponse(
            name=path.name,
            type="TXT",
            mode="text",
            content=path.read_text(encoding="utf-8", errors="ignore")[:12000],
        )

    if extension == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        html = "".join(f"<p>{escape_html(paragraph)}</p>" for paragraph in paragraphs[:120])
        return DocumentPreviewResponse(
            name=path.name,
            type="DOCX",
            mode="html",
            content=html or "<p>Aucun texte lisible dans ce document.</p>",
        )

    raise HTTPException(status_code=400, detail="Format non supporte pour la previsualisation.")


@app.get("/documents/{filename}/metadata", response_model=DocumentMetadataResponse)
def get_document_metadata(filename: str, current_user: dict = Depends(optional_current_user)):
    """Retourne les metadonnees GED d'un document."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    metadata = get_metadata_for(path.name)
    return DocumentMetadataResponse(filename=path.name, **metadata)


@app.put("/documents/{filename}/metadata", response_model=DocumentMetadataResponse)
def update_document_metadata(
    filename: str,
    payload: DocumentMetadata,
    current_user: dict = Depends(require_permission("documents:write")),
):
    """Met a jour les metadonnees GED d'un document."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    all_metadata = load_metadata_store()
    previous_metadata = get_metadata_for(path.name)
    clean_payload = normalize_metadata(payload)
    previous_metadata.update(clean_payload)
    clean_payload = previous_metadata
    clean_payload["updated_at"] = datetime.now().isoformat(timespec="seconds")
    all_metadata[path.name] = clean_payload
    save_metadata_store(all_metadata)
    return DocumentMetadataResponse(filename=path.name, **clean_payload)


@app.get("/documents/{filename}/workflow", response_model=WorkflowHistoryResponse)
def get_document_workflow(filename: str, current_user: dict = Depends(optional_current_user)):
    """Retourne l'historique de workflow d'un document."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    metadata = get_metadata_for(path.name)
    actions = [
        WorkflowActionItem(**item)
        for item in load_workflow_store().get(path.name, [])
    ]
    return WorkflowHistoryResponse(
        filename=path.name,
        current_status=metadata.get("status", "Indexe"),
        steps=WORKFLOW_STEPS,
        actions=actions,
    )


@app.post("/documents/{filename}/workflow", response_model=WorkflowHistoryResponse)
def transition_document_workflow(
    filename: str,
    payload: WorkflowTransitionRequest,
    current_user: dict = Depends(require_permission("workflow:comment")),
):
    """Change l'etat documentaire et ajoute une entree d'historique."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    target_status = normalize_workflow_status(payload.status)
    current_metadata = get_metadata_for(path.name)
    previous_status = current_metadata.get("status", "Indexe")

    all_metadata = load_metadata_store()
    current_metadata["status"] = target_status
    current_metadata["updated_at"] = datetime.now().isoformat(timespec="seconds")
    all_metadata[path.name] = current_metadata
    save_metadata_store(all_metadata)

    workflow_store = load_workflow_store()
    action = {
        "filename": path.name,
        "from_status": previous_status,
        "to_status": target_status,
        "comment": payload.comment.strip(),
        "actor": payload.actor.strip() or "Utilisateur demo",
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    workflow_store.setdefault(path.name, []).append(action)
    save_workflow_store(workflow_store)

    return get_document_workflow(path.name)


@app.get("/metadata/categories")
def list_categories():
    """Liste les categories utilisees dans la bibliotheque documentaire."""
    categories = {
        get_metadata_for(document.name)["category"]
        for document in discover_documents()
        if get_metadata_for(document.name)["category"]
    }
    return {"categories": sorted(categories)}


@app.delete("/documents/{filename}")
def delete_document(
    filename: str,
    actor: str = "Utilisateur demo",
    current_user: dict = Depends(require_permission("documents:delete")),
):
    """Place un document dans la corbeille sans supprimer le fichier original."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    metadata_store = load_metadata_store()
    metadata = get_metadata_for(path.name)
    metadata["deleted"] = True
    metadata["deleted_at"] = datetime.now().isoformat(timespec="seconds")
    metadata["deleted_by"] = actor.strip() or "Utilisateur demo"
    metadata["updated_at"] = metadata["deleted_at"]
    metadata_store[path.name] = metadata
    save_metadata_store(metadata_store)
    append_workflow_action(path.name, metadata.get("status", "Indexe"), "Corbeille", "Suppression logique", actor)
    return {"status": "deleted", "filename": path.name}


@app.post("/documents/{filename}/restore")
def restore_document(
    filename: str,
    actor: str = "Utilisateur demo",
    current_user: dict = Depends(require_permission("documents:delete")),
):
    """Restaure un document place dans la corbeille."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    metadata_store = load_metadata_store()
    metadata = get_metadata_for(path.name)
    metadata["deleted"] = False
    metadata["deleted_at"] = None
    metadata["deleted_by"] = ""
    metadata["updated_at"] = datetime.now().isoformat(timespec="seconds")
    metadata_store[path.name] = metadata
    save_metadata_store(metadata_store)
    append_workflow_action(path.name, "Corbeille", metadata.get("status", "Indexe"), "Restauration", actor)
    return {"status": "restored", "filename": path.name}


@app.get("/documents/{filename}/versions", response_model=DocumentVersionsResponse)
def get_document_versions(filename: str, current_user: dict = Depends(optional_current_user)):
    """Retourne les versions connues d'un document."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    versions = get_versions_for(path)
    return DocumentVersionsResponse(filename=path.name, versions=versions)


@app.get("/documents/{filename}/similar", response_model=SimilarDocumentsResponse)
def get_similar_documents(
    filename: str,
    limit: int = Query(5, ge=1, le=10),
    current_user: dict = Depends(optional_current_user),
):
    """Detecte les documents proches ou les doublons potentiels."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    items = find_similar_documents(path, limit=limit)
    return SimilarDocumentsResponse(filename=path.name, items=items)


@app.get("/documents/{filename}/ai", response_model=DocumentAIAnalysisResponse)
def analyze_document(filename: str, current_user: dict = Depends(optional_current_user)):
    """Produit une analyse IA locale : resume, mots-cles, categorie et doublons."""
    path = resolve_raw_document(filename)
    ensure_document_access(path.name, current_user)
    text = read_document_text(path)
    keywords = extract_keywords(text)
    return DocumentAIAnalysisResponse(
        filename=path.name,
        summary=summarize_text(text),
        keywords=keywords,
        suggested_category=suggest_category(text, keywords),
        confidentiality=estimate_confidentiality(text, keywords),
        similar_documents=find_similar_documents(path, limit=3),
    )


@app.post("/assistant/ask", response_model=AssistantAnswerResponse)
def ask_document_assistant(payload: AssistantQuestionRequest):
    """Assistant documentaire RAG local base sur la recherche hybride et les sources."""
    search_payload = SearchRequest(
        query=payload.question,
        top_k=payload.top_k,
        mode="hybrid",
        semantic_weight=0.7,
        keyword_weight=0.3,
    )
    search_response = search_documents(search_payload)
    sources = [
        AssistantSourceItem(
            filename=result.filename,
            chunk_id=result.chunk_id,
            score=result.score,
            text_preview=result.text_preview,
        )
        for result in search_response.results
    ]
    answer = build_assistant_answer(payload.question, sources)
    confidence = round(max((source.score for source in sources), default=0.0), 4)
    return AssistantAnswerResponse(
        question=payload.question,
        answer=answer,
        confidence=confidence,
        sources=sources,
    )


@app.get("/users", response_model=list[UserItem])
def list_users(current_user: dict = Depends(require_permission("users:read"))):
    """Liste les utilisateurs et roles de demonstration."""
    return [UserItem(**user) for user in load_users_store()]


@app.post("/auth/login", response_model=LoginResponse)
def login_demo(payload: LoginRequest):
    """Connexion JWT de demonstration."""
    email = payload.email.strip().lower()
    for user in load_users_store():
        if user["email"].lower() == email and verify_password(payload.password, user.get("password_hash", "")):
            token = create_jwt(user)
            return LoginResponse(token=token, user=UserItem(**user))
    raise HTTPException(status_code=401, detail="Identifiants invalides.")


@app.post("/upload")
async def upload_document(
    request: Request,
    filename: str = Query(..., min_length=1, description="Nom du fichier a importer"),
    current_user: dict = Depends(require_permission("documents:upload")),
):
    """
    Importe un document, extrait son texte et reconstruit l'index vectoriel.

    Le fichier est envoye en binaire brut afin d'eviter une dependance
    supplementaire a python-multipart.
    """
    safe_filename = sanitize_filename(filename)
    extension = Path(safe_filename).suffix.lower()

    if extension not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Format non supporte. Formats acceptes : PDF, DOCX, TXT.",
        )

    content = await request.body()
    if not content:
        raise HTTPException(status_code=400, detail="Fichier vide.")

    raw_path = unique_storage_path(Path(RAW_DIR), safe_filename)
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    raw_path.write_bytes(content)

    try:
        processed_path = ingest_uploaded_file(raw_path)
        initialize_metadata_for(raw_path.name)
        register_document_version(raw_path)
        rebuild_index()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Echec de l'indexation : {exc}") from exc

    return {
        "status": "indexed",
        "filename": raw_path.name,
        "raw_path": str(raw_path),
        "processed_path": str(processed_path),
        "stats": stats(),
    }


def ingest_uploaded_file(raw_path: Path) -> Path:
    """Extrait le texte d'un fichier importe."""
    from src.ingestion.loader import process_file

    return process_file(raw_path, PROCESSED_DIR)


def rebuild_index() -> None:
    """Reconstruit l'index FAISS apres ajout d'un document."""
    from src.embeddings.build_index import build_index

    build_index(PROCESSED_DIR, EMBEDDINGS_DIR, allow_model_download=False)


def sanitize_filename(filename: str) -> str:
    """Conserve un nom de fichier simple et empeche les chemins relatifs."""
    name = Path(filename).name.strip()
    name = re.sub(r"[^A-Za-z0-9_.-]+", "_", name)
    if not name or name in {".", ".."}:
        raise HTTPException(status_code=400, detail="Nom de fichier invalide.")
    return name


def unique_storage_path(directory: Path, filename: str) -> Path:
    """Cree un chemin de stockage sans ecraser un fichier deja importe."""
    candidate = directory / filename
    if not candidate.exists():
        return candidate

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = candidate.stem
    suffix = candidate.suffix
    return directory / f"{stem}_{timestamp}{suffix}"


def discover_documents(include_deleted: bool = False, current_user: dict | None = None) -> list[DocumentItem]:
    """Construit la liste des documents depuis data/raw."""
    raw_path = Path(RAW_DIR)
    if not raw_path.exists():
        return []

    metadata_store = load_metadata_store()
    documents = []
    for file in sorted(raw_path.iterdir()):
        if not file.is_file() or file.suffix.lower() not in SUPPORTED_UPLOAD_EXTENSIONS:
            continue

        stat = file.stat()
        extension = file.suffix.lower()
        metadata = metadata_store.get(file.name, default_metadata())
        if metadata.get("deleted") and not include_deleted:
            continue
        if current_user:
            allowed_roles = metadata.get("allowed_roles", [])
            if allowed_roles and current_user.get("role") not in allowed_roles and not has_permission(current_user, "documents:*"):
                continue
        documents.append(
            DocumentItem(
                name=file.name,
                title=file.stem.replace("_", " "),
                extension=extension,
                type=extension.lstrip(".").upper(),
                size_bytes=stat.st_size,
                size_label=format_size(stat.st_size),
                added_at=datetime.fromtimestamp(stat.st_ctime).isoformat(timespec="seconds"),
                modified_at=datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                page_count=count_pages(file),
                status=metadata.get("status") or ("Indexe" if processed_file_exists(file) else "Importe"),
                processed=processed_file_exists(file),
                category=metadata.get("category", "Non classe"),
                author=metadata.get("author", ""),
                description=metadata.get("description", ""),
                tags=metadata.get("tags", []),
                folder=metadata.get("folder", "/"),
                allowed_roles=metadata.get("allowed_roles", []),
                deleted=bool(metadata.get("deleted", False)),
                deleted_at=metadata.get("deleted_at"),
                download_url=f"/documents/{file.name}/download",
                preview_url=f"/documents/{file.name}/preview",
            )
        )
    return documents


def sort_documents(documents: list[DocumentItem], sort_by: str) -> list[DocumentItem]:
    """Trie les documents selon un critere attendu par l'interface."""
    if sort_by == "name_asc":
        return sorted(documents, key=lambda doc: doc.title.lower())
    if sort_by == "name_desc":
        return sorted(documents, key=lambda doc: doc.title.lower(), reverse=True)
    if sort_by == "size_desc":
        return sorted(documents, key=lambda doc: doc.size_bytes, reverse=True)
    if sort_by == "type_asc":
        return sorted(documents, key=lambda doc: (doc.type, doc.title.lower()))
    return sorted(documents, key=lambda doc: doc.modified_at, reverse=True)


def resolve_raw_document(filename: str) -> Path:
    """Resout un document brut sans permettre de sortie du dossier data/raw."""
    safe_name = sanitize_filename(filename)
    path = Path(RAW_DIR) / safe_name
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=404, detail="Document introuvable.")
    if path.suffix.lower() not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Format non supporte.")
    return path


def processed_file_exists(raw_file: Path) -> bool:
    processed_name = f"{raw_file.stem}_{raw_file.suffix.lstrip('.')}.txt"
    return (Path(PROCESSED_DIR) / processed_name).exists()


def processed_text_path(raw_file: Path) -> Path:
    processed_name = f"{raw_file.stem}_{raw_file.suffix.lstrip('.')}.txt"
    return Path(PROCESSED_DIR) / processed_name


def read_document_text(raw_file: Path) -> str:
    """Lit le texte extrait quand il existe, sinon effectue une lecture legere."""
    processed_path = processed_text_path(raw_file)
    if processed_path.exists():
        return processed_path.read_text(encoding="utf-8", errors="ignore")

    if raw_file.suffix.lower() == ".txt":
        return raw_file.read_text(encoding="utf-8", errors="ignore")

    if raw_file.suffix.lower() == ".docx":
        try:
            from docx import Document

            document = Document(str(raw_file))
            return "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        except Exception:
            return ""

    return ""


def count_pages(path: Path) -> int | None:
    """Retourne le nombre de pages PDF ou None pour les autres formats."""
    if path.suffix.lower() != ".pdf":
        return None
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path)).pages)
    except Exception:
        return None


def format_size(size: int) -> str:
    if size < 1024:
        return f"{size} o"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} Ko"
    return f"{size / (1024 * 1024):.1f} Mo"


def escape_html(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def load_metadata_store() -> dict:
    """Charge les metadonnees depuis la base relationnelle."""
    init_db(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        rows = rows_to_dicts(execute(connection, "SELECT * FROM document_metadata").fetchall())
    metadata = {}
    for row in rows:
        metadata[row["filename"]] = metadata_row_to_dict(row)
    return metadata


def migrate_legacy_json_files() -> None:
    """Importe une seule fois les anciens fichiers JSON vers la base SQL."""
    init_db(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        metadata_count = execute(connection, "SELECT COUNT(*) AS count FROM document_metadata").fetchone()["count"]
        workflow_count = execute(connection, "SELECT COUNT(*) AS count FROM workflow_actions").fetchone()["count"]
        versions_count = execute(connection, "SELECT COUNT(*) AS count FROM document_versions").fetchone()["count"]

    metadata_path = Path(METADATA_PATH)
    workflow_path = Path(WORKFLOW_LOG_PATH)
    versions_path = Path(VERSIONS_PATH)

    if metadata_count == 0 and metadata_path.exists():
        try:
            legacy_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            if isinstance(legacy_metadata, dict):
                save_metadata_store(legacy_metadata)
        except json.JSONDecodeError:
            pass

    if workflow_count == 0 and workflow_path.exists():
        try:
            legacy_workflow = json.loads(workflow_path.read_text(encoding="utf-8"))
            if isinstance(legacy_workflow, dict):
                save_workflow_store(legacy_workflow)
        except json.JSONDecodeError:
            pass

    if versions_count == 0 and versions_path.exists():
        try:
            legacy_versions = json.loads(versions_path.read_text(encoding="utf-8"))
            if isinstance(legacy_versions, dict):
                save_versions_store(legacy_versions)
        except json.JSONDecodeError:
            pass


def save_metadata_store(metadata: dict) -> None:
    """Sauvegarde les metadonnees GED dans la base relationnelle."""
    init_db(METADATA_PATH)
    mark = placeholder(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        for filename, values in metadata.items():
            merged = default_metadata()
            merged.update(values)
            execute(
                connection,
                f"""
                INSERT INTO document_metadata (
                    filename, category, author, description, tags, status, folder,
                    allowed_roles, updated_at, deleted, deleted_at, deleted_by
                )
                VALUES ({','.join([mark] * 12)})
                ON CONFLICT(filename) DO UPDATE SET
                    category=excluded.category,
                    author=excluded.author,
                    description=excluded.description,
                    tags=excluded.tags,
                    status=excluded.status,
                    folder=excluded.folder,
                    allowed_roles=excluded.allowed_roles,
                    updated_at=excluded.updated_at,
                    deleted=excluded.deleted,
                    deleted_at=excluded.deleted_at,
                    deleted_by=excluded.deleted_by
                """,
                (
                    filename,
                    merged["category"],
                    merged["author"],
                    merged["description"],
                    encode_json(merged["tags"]),
                    merged["status"],
                    merged["folder"],
                    encode_json(merged["allowed_roles"]),
                    merged["updated_at"],
                    bool(merged["deleted"]),
                    merged["deleted_at"],
                    merged["deleted_by"],
                ),
            )


def default_metadata() -> dict:
    return {
        "category": "Non classe",
        "author": "",
        "description": "",
        "tags": [],
        "status": "Indexe",
        "folder": "/",
        "allowed_roles": [],
        "updated_at": None,
        "deleted": False,
        "deleted_at": None,
        "deleted_by": "",
    }


def get_metadata_for(filename: str) -> dict:
    metadata = load_metadata_store().get(filename, {})
    merged = default_metadata()
    merged.update(metadata)
    return merged


def metadata_row_to_dict(row: dict) -> dict:
    return {
        "category": row.get("category") or "Non classe",
        "author": row.get("author") or "",
        "description": row.get("description") or "",
        "tags": decode_json(row.get("tags"), []),
        "status": row.get("status") or "Indexe",
        "folder": row.get("folder") or "/",
        "allowed_roles": decode_json(row.get("allowed_roles"), []),
        "updated_at": row.get("updated_at"),
        "deleted": bool(row.get("deleted")),
        "deleted_at": row.get("deleted_at"),
        "deleted_by": row.get("deleted_by") or "",
    }


def initialize_metadata_for(filename: str) -> None:
    metadata = load_metadata_store()
    metadata.setdefault(filename, default_metadata())
    save_metadata_store(metadata)
    workflow_store = load_workflow_store()
    workflow_store.setdefault(filename, [])
    save_workflow_store(workflow_store)


def append_workflow_action(
    filename: str,
    from_status: str,
    to_status: str,
    comment: str = "",
    actor: str = "Utilisateur demo",
) -> None:
    workflow_store = load_workflow_store()
    workflow_store.setdefault(filename, []).append(
        {
            "filename": filename,
            "from_status": from_status,
            "to_status": to_status,
            "comment": comment,
            "actor": actor.strip() or "Utilisateur demo",
            "created_at": datetime.now().isoformat(timespec="seconds"),
        }
    )
    save_workflow_store(workflow_store)


def normalize_metadata(payload: DocumentMetadata) -> dict:
    tags = []
    seen = set()
    for tag in payload.tags:
        cleaned = tag.strip()
        key = cleaned.lower()
        if cleaned and key not in seen:
            seen.add(key)
            tags.append(cleaned[:50])

    return {
        "category": payload.category.strip() or "Non classe",
        "author": payload.author.strip(),
        "description": payload.description.strip(),
        "tags": tags[:12],
        "status": payload.status.strip() or "Indexe",
        "folder": normalize_folder(payload.folder),
        "allowed_roles": normalize_roles(payload.allowed_roles),
    }


def normalize_folder(folder: str) -> str:
    cleaned = re.sub(r"/+", "/", f"/{folder.strip().strip('/')}")
    return cleaned[:160] or "/"


def normalize_roles(roles: list[str]) -> list[str]:
    allowed = {"Administrateur", "Responsable qualite", "Validateur", "Employe"}
    normalized = []
    for role in roles:
        cleaned = role.strip()
        if cleaned in allowed and cleaned not in normalized:
            normalized.append(cleaned)
    return normalized


def normalize_workflow_status(status: str) -> str:
    """Valide un statut du workflow documentaire."""
    cleaned = status.strip()
    for allowed in WORKFLOW_STEPS:
        if cleaned.lower() == allowed.lower():
            return allowed
    raise HTTPException(
        status_code=400,
        detail=f"Statut invalide. Valeurs autorisees : {', '.join(WORKFLOW_STEPS)}",
    )


def load_workflow_store() -> dict:
    """Charge l'historique des workflows depuis la base relationnelle."""
    init_db(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        rows = rows_to_dicts(
            execute(
                connection,
                "SELECT filename, from_status, to_status, comment, actor, created_at FROM workflow_actions ORDER BY created_at ASC",
            ).fetchall()
        )
    workflow: dict[str, list[dict]] = {}
    for row in rows:
        workflow.setdefault(row["filename"], []).append(row)
    return workflow


def save_workflow_store(workflow: dict) -> None:
    """Sauvegarde l'historique des workflows dans la base relationnelle."""
    init_db(METADATA_PATH)
    mark = placeholder(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        execute(connection, "DELETE FROM workflow_actions")
        for filename, actions in workflow.items():
            for action in actions:
                execute(
                    connection,
                    f"""
                    INSERT INTO workflow_actions
                    (filename, from_status, to_status, comment, actor, created_at)
                    VALUES ({','.join([mark] * 6)})
                    """,
                    (
                        filename,
                        action.get("from_status", ""),
                        action.get("to_status", ""),
                        action.get("comment", ""),
                        action.get("actor", "Utilisateur demo"),
                        action.get("created_at", datetime.now().isoformat(timespec="seconds")),
                    ),
                )


def count_buckets(values) -> list[DashboardBucket]:
    """Compte une sequence de valeurs pour alimenter les graphiques."""
    counts: dict[str, int] = {}
    for value in values:
        label = str(value or "Non renseigne")
        counts[label] = counts.get(label, 0) + 1
    return [
        DashboardBucket(label=label, count=count)
        for label, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    ]


def collect_recent_workflow_actions(limit: int = 8) -> list[WorkflowActionItem]:
    """Retourne les dernieres actions de workflow, tous documents confondus."""
    actions = []
    for entries in load_workflow_store().values():
        actions.extend(entries)

    actions.sort(key=lambda item: item.get("created_at", ""), reverse=True)
    return [WorkflowActionItem(**item) for item in actions[:limit]]


def load_versions_store() -> dict:
    init_db(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        rows = rows_to_dicts(
            execute(
                connection,
                "SELECT filename, version, checksum, size_bytes, size_label, created_at, source FROM document_versions ORDER BY filename, version",
            ).fetchall()
        )
    versions: dict[str, list[dict]] = {}
    for row in rows:
        versions.setdefault(row["filename"], []).append(row)
    return versions


def save_versions_store(versions: dict) -> None:
    init_db(METADATA_PATH)
    mark = placeholder(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        execute(connection, "DELETE FROM document_versions")
        for filename, items in versions.items():
            for item in items:
                execute(
                    connection,
                    f"""
                    INSERT INTO document_versions
                    (filename, version, checksum, size_bytes, size_label, created_at, source)
                    VALUES ({','.join([mark] * 7)})
                    ON CONFLICT(filename, checksum) DO NOTHING
                    """,
                    (
                        filename,
                        item["version"],
                        item["checksum"],
                        item["size_bytes"],
                        item["size_label"],
                        item["created_at"],
                        item.get("source", "upload"),
                    ),
                )


def file_checksum(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def register_document_version(path: Path) -> None:
    versions = load_versions_store()
    existing_versions = versions.setdefault(path.name, [])
    checksum = file_checksum(path)
    if any(item.get("checksum") == checksum for item in existing_versions):
        return

    existing_versions.append(
        {
            "filename": path.name,
            "version": len(existing_versions) + 1,
            "checksum": checksum,
            "size_bytes": path.stat().st_size,
            "size_label": format_size(path.stat().st_size),
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "source": "upload",
        }
    )
    save_versions_store(versions)


def get_versions_for(path: Path) -> list[DocumentVersionItem]:
    versions = load_versions_store().get(path.name, [])
    if not versions:
        versions = [
            {
                "filename": path.name,
                "version": 1,
                "checksum": file_checksum(path),
                "size_bytes": path.stat().st_size,
                "size_label": format_size(path.stat().st_size),
                "created_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
                "source": "existing-file",
            }
        ]
    return [DocumentVersionItem(**item) for item in versions]


def tokenize(text: str) -> list[str]:
    return [
        token
        for token in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]{3,}", text.lower())
        if token not in FRENCH_STOPWORDS
    ]


def extract_keywords(text: str, limit: int = 10) -> list[str]:
    counts: dict[str, int] = {}
    for token in tokenize(text):
        counts[token] = counts.get(token, 0) + 1
    return [
        token
        for token, _count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:limit]
    ]


def summarize_text(text: str, max_sentences: int = 3) -> str:
    cleaned = " ".join(text.split())
    if not cleaned:
        return "Aucun texte exploitable n'a ete trouve pour produire un resume."

    sentences = re.split(r"(?<=[.!?])\s+", cleaned)
    sentences = [sentence.strip() for sentence in sentences if len(sentence.strip()) > 30]
    if not sentences:
        return cleaned[:500]

    keywords = set(extract_keywords(cleaned, limit=12))
    ranked = sorted(
        enumerate(sentences[:40]),
        key=lambda item: sum(1 for token in tokenize(item[1]) if token in keywords),
        reverse=True,
    )
    selected_indexes = sorted(index for index, _sentence in ranked[:max_sentences])
    return " ".join(sentences[index] for index in selected_indexes)[:900]


def suggest_category(text: str, keywords: list[str]) -> str:
    tokens = set(tokenize(text))
    tokens.update(keyword.lower() for keyword in keywords)
    scores = {
        category: len(tokens.intersection(rule_tokens))
        for category, rule_tokens in CATEGORY_RULES.items()
    }
    best_category, best_score = max(scores.items(), key=lambda item: item[1])
    return best_category if best_score > 0 else "Non classe"


def estimate_confidentiality(text: str, keywords: list[str]) -> str:
    tokens = set(tokenize(text))
    tokens.update(keyword.lower() for keyword in keywords)
    if tokens.intersection({"salaire", "paie", "confidentiel", "contrat", "password", "motdepasse"}):
        return "Confidentiel"
    if tokens.intersection({"budget", "finance", "facture", "rh", "personnel"}):
        return "Interne"
    return "Public interne"


def cosine_from_tokens(left_text: str, right_text: str) -> float:
    left_counts: dict[str, int] = {}
    right_counts: dict[str, int] = {}
    for token in tokenize(left_text):
        left_counts[token] = left_counts.get(token, 0) + 1
    for token in tokenize(right_text):
        right_counts[token] = right_counts.get(token, 0) + 1

    if not left_counts or not right_counts:
        return 0.0

    common = set(left_counts).intersection(right_counts)
    numerator = sum(left_counts[token] * right_counts[token] for token in common)
    left_norm = sum(value * value for value in left_counts.values()) ** 0.5
    right_norm = sum(value * value for value in right_counts.values()) ** 0.5
    return numerator / (left_norm * right_norm) if left_norm and right_norm else 0.0


def find_similar_documents(path: Path, limit: int = 5) -> list[SimilarDocumentItem]:
    base_text = read_document_text(path)
    items: list[SimilarDocumentItem] = []
    for document in discover_documents():
        if document.name == path.name:
            continue
        candidate_path = Path(RAW_DIR) / document.name
        similarity = cosine_from_tokens(base_text, read_document_text(candidate_path))
        if similarity <= 0:
            continue
        if similarity >= 0.85:
            reason = "Doublon potentiel"
        elif similarity >= 0.45:
            reason = "Document tres proche"
        else:
            reason = "Sujet lie"
        items.append(
            SimilarDocumentItem(
                filename=document.name,
                title=document.title,
                type=document.type,
                similarity=round(similarity, 4),
                reason=reason,
            )
        )
    return sorted(items, key=lambda item: item.similarity, reverse=True)[:limit]


def build_assistant_answer(question: str, sources: list[AssistantSourceItem]) -> str:
    if not sources:
        return "Je n'ai pas trouve de passage suffisamment pertinent dans les documents indexes."

    question_tokens = set(tokenize(question))
    candidate_sentences = []
    for source in sources[:4]:
        sentences = re.split(r"(?<=[.!?])\s+", source.text_preview)
        for sentence in sentences:
            score = sum(1 for token in tokenize(sentence) if token in question_tokens)
            if len(sentence.strip()) > 25:
                candidate_sentences.append((score, source.score, sentence.strip(), source.filename))

    candidate_sentences.sort(key=lambda item: (item[0], item[1]), reverse=True)
    selected = candidate_sentences[:3]
    if not selected:
        return f"Les passages les plus proches indiquent notamment : {sources[0].text_preview[:500]}"

    answer = " ".join(sentence for _score, _source_score, sentence, _filename in selected)
    filenames = sorted({filename for _score, _source_score, _sentence, filename in selected})
    return f"{answer[:900]} Sources principales : {', '.join(filenames)}."


def default_users() -> list[dict]:
    return [
        {
            "id": "admin",
            "name": "Administrateur GED",
            "email": "admin@inrh.demo",
            "role": "Administrateur",
            "permissions": ["documents:*", "users:*", "workflow:*", "audit:read"],
            "password_hash": hash_password("demo123"),
        },
        {
            "id": "qualite",
            "name": "Responsable qualite",
            "email": "qualite@inrh.demo",
            "role": "Responsable qualite",
            "permissions": ["documents:read", "documents:write", "documents:validate", "workflow:*", "audit:read"],
            "password_hash": hash_password("demo123"),
        },
        {
            "id": "validateur",
            "name": "Validateur documentaire",
            "email": "validateur@inrh.demo",
            "role": "Validateur",
            "permissions": ["documents:read", "documents:validate", "workflow:comment"],
            "password_hash": hash_password("demo123"),
        },
        {
            "id": "employe",
            "name": "Employe INRH",
            "email": "employe@inrh.demo",
            "role": "Employe",
            "permissions": ["documents:read", "documents:upload", "workflow:submit"],
            "password_hash": hash_password("demo123"),
        },
    ]


def load_users_store() -> list[dict]:
    init_db(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        rows = rows_to_dicts(
            execute(
                connection,
                "SELECT id, name, email, role, permissions, password_hash FROM users WHERE is_active = 1",
            ).fetchall()
        )
    users = []
    for row in rows:
        row["permissions"] = decode_json(row.get("permissions"), [])
        users.append(row)
    return users


def seed_default_users() -> None:
    init_db(METADATA_PATH)
    mark = placeholder(METADATA_PATH)
    with connect(METADATA_PATH) as connection:
        for user in default_users():
            execute(
                connection,
                f"""
                INSERT INTO users (id, name, email, role, permissions, password_hash, is_active)
                VALUES ({','.join([mark] * 7)})
                ON CONFLICT(id) DO NOTHING
                """,
                (
                    user["id"],
                    user["name"],
                    user["email"],
                    user["role"],
                    encode_json(user["permissions"]),
                    user["password_hash"],
                    True,
                ),
            )


def hash_password(password: str) -> str:
    salt = os.getenv("GED_PASSWORD_SALT", "ged-demo-salt")
    return sha256(f"{salt}|{password}".encode("utf-8")).hexdigest()


def verify_password(password: str, password_hash: str) -> bool:
    return hmac.compare_digest(hash_password(password), password_hash)
